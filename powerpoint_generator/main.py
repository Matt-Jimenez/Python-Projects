"""
PowerPoint presentation generator.

Uses python‑pptx to create a presentation with a title slide and a chart slide based on data from a CSV file.
"""

import pandas as pd
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def generate_presentation():
    """Build a PowerPoint file from data.csv."""
    df = pd.read_csv('data.csv')
    prs = Presentation()
    # Title slide
    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = 'Report'
    slide.placeholders[1].text = 'Generated presentation'
    # Chart slide
    slide_layout = prs.slide_layouts[5]
    chart_slide = prs.slides.add_slide(slide_layout)
    chart_slide.shapes.title.text = 'Data Chart'
    chart_data = ChartData()
    chart_data.categories = df.iloc[:, 0].tolist()
    chart_data.add_series('Series 1', df.iloc[:, 1].tolist())
    x, y, cx, cy = Inches(1), Inches(1.5), Inches(8), Inches(4)
    chart = chart_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data
    ).chart
    prs.save('presentation.pptx')


if __name__ == '__main__':
    generate_presentation()