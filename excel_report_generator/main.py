"""
Excel report generator.

Reads a CSV file into a pandas DataFrame, computes summary statistics, writes them to an Excel workbook using openpyxl, and adds a chart.
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference


def generate_report():
    """Generate an Excel report from data.csv."""
    df = pd.read_csv('data.csv')
    summary = df.describe()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    # Header row
    ws.append(['Statistic'] + list(summary.columns))
    # Data rows
    for stat in summary.index:
        ws.append([stat] + list(summary.loc[stat]))
    # Create a bar chart
    chart = BarChart()
    data = Reference(ws, min_col=2, min_row=2, max_row=ws.max_row, max_col=ws.max_column)
    categories = Reference(ws, min_col=1, min_row=3, max_row=ws.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.title = 'Summary Chart'
    ws.add_chart(chart, 'H2')
    wb.save('report.xlsx')


if __name__ == '__main__':
    generate_report()