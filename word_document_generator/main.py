"""
Word document generator.

Uses python‑docx to create a sample document with headings, paragraphs, a table, and saves it as a .docx file.
"""

from docx import Document


def create_document():
    """Create a sample Word document."""
    doc = Document()
    doc.add_heading('Sample Document', level=1)
    doc.add_paragraph('This is a sample document created with python‑docx.')
    doc.add_paragraph('You can modify this script to add tables and images.')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Header 1'
    table.cell(0, 1).text = 'Header 2'
    table.cell(1, 0).text = 'Data 1'
    table.cell(1, 1).text = 'Data 2'
    doc.save('document.docx')


if __name__ == '__main__':
    create_document()